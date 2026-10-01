"""
Automated unit tests for multi-format document text extraction engine (TASK-17).
Validates .txt, .docx, and .pdf parsing, parser fallback, and fault tolerance.
"""

from pathlib import Path
import docx
from fpdf import FPDF
import pytest

from src.extractor import (
    extract_from_txt,
    extract_from_docx,
    extract_from_pdf,
    extract_text,
)
from src import config

def test_extract_txt_utf8(tmp_path: Path):
    """Verify clean extraction from UTF-8 encoded plain text file."""
    sample_text = "Senior Python Developer with expertise in FastAPI and SQL."
    txt_file = tmp_path / "sample_utf8.txt"
    txt_file.write_text(sample_text, encoding="utf-8")

    result = extract_from_txt(txt_file)
    assert sample_text == result
    assert extract_text(txt_file) == sample_text

def test_extract_txt_latin1_fallback(tmp_path: Path):
    """Verify fallback to Latin-1 encoding for non-UTF8 text files."""
    sample_text = "Développeur Python avec expérience en modélisation"
    txt_file = tmp_path / "sample_latin1.txt"
    with open(txt_file, "wb") as f:
        f.write(sample_text.encode("latin-1"))

    result = extract_from_txt(txt_file)
    assert len(result) > 0
    assert "Python" in result

def test_extract_docx_paragraphs_and_tables(tmp_path: Path):
    """Verify extraction of body paragraphs and table cell contents from .docx."""
    doc_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_paragraph("Alex Morgan - Senior Engineer")
    doc.add_paragraph("Skills: Python, Docker, Kubernetes")

    # Add a table with skill details
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Category"
    table.rows[0].cells[1].text = "Skill"
    table.rows[1].cells[0].text = "Backend"
    table.rows[1].cells[1].text = "FastAPI"

    doc.save(str(doc_path))

    result = extract_from_docx(doc_path)
    assert "Alex Morgan - Senior Engineer" in result
    assert "Skills: Python, Docker, Kubernetes" in result
    assert "Backend | FastAPI" in result
    assert extract_text(doc_path) == result

def test_extract_pdf_generation_and_reading(tmp_path: Path):
    """Verify PDF parsing using a programmatically generated test PDF."""
    pdf_path = tmp_path / "sample.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(w=0, h=10, text="Marcus Vance - Backend Infrastructure Specialist", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(w=0, h=10, text="Technologies: PostgreSQL, Redis, Linux, Docker", new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(pdf_path))

    result = extract_from_pdf(pdf_path)
    assert "Marcus Vance" in result
    assert "PostgreSQL" in result
    assert extract_text(pdf_path) == result

def test_extract_edge_cases(tmp_path: Path):
    """Verify graceful handling for non-existent, empty, unsupported, and corrupted files."""
    # 1. Non-existent file
    non_existent = tmp_path / "does_not_exist.pdf"
    assert extract_text(non_existent) == ""

    # 2. Empty 0-byte file
    empty_file = tmp_path / "empty.pdf"
    empty_file.touch()
    assert extract_text(empty_file) == ""

    # 3. Unsupported extension
    unsupported = tmp_path / "data.xyz"
    unsupported.write_text("Some text", encoding="utf-8")
    assert extract_text(unsupported) == ""

    # 4. Corrupted PDF structure
    corrupt_pdf = tmp_path / "corrupt.pdf"
    corrupt_pdf.write_bytes(b"%PDF-1.4\nCorrupted binary noise instead of valid objects\n%%EOF")
    # Should safely return empty string without unhandled exception
    assert extract_text(corrupt_pdf) == ""

    # 5. Corrupted DOCX structure
    corrupt_docx = tmp_path / "corrupt.docx"
    corrupt_docx.write_bytes(b"Not a valid zip or docx archive")
    assert extract_text(corrupt_docx) == ""

def test_extract_all_10_curated_resumes():
    """Verify that all 10 actual curated resumes in resumes/ extract substantive text."""
    resume_files = [
        rf for rf in config.RESUMES_DIR.iterdir()
        if rf.is_file() and rf.suffix.lower() in config.SUPPORTED_EXTENSIONS
    ]
    assert len(resume_files) == 10

    for rf in resume_files:
        text = extract_text(rf)
        assert isinstance(text, str)
        assert len(text.strip()) >= 100, f"Resume {rf.name} extracted insufficient text ({len(text)} chars)"
