"""
Automated Resume Screening Tool — Document Extraction Engine.

Provides robust, fault-tolerant extraction of textual content across .pdf,
.docx, and .txt files with automatic parser fallbacks and comprehensive error trapping.
"""

import logging
from pathlib import Path
import pdfplumber
import pypdf
import docx

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

def extract_from_txt(file_path: Path) -> str:
    """
    Extract plain text from a .txt file with encoding fallback.

    Args:
        file_path: Path to the .txt file.

    Returns:
        str: Extracted textual content.
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read().strip()
    except UnicodeDecodeError:
        logger.warning("UTF-8 decode failed for %s. Falling back to Latin-1.", file_path.name)
        with open(file_path, "r", encoding="latin-1", errors="replace") as f:
            return f.read().strip()
    except Exception as e:
        logger.error("Error reading text file %s: %s", file_path.name, e)
        return ""

def extract_from_docx(file_path: Path) -> str:
    """
    Extract text from a Microsoft Word (.docx) document including body
    paragraphs and embedded table cells.

    Args:
        file_path: Path to the .docx file.

    Returns:
        str: Extracted textual content.
    """
    try:
        doc = docx.Document(file_path)
        extracted_lines = []

        # Extract main body paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                extracted_lines.append(text)

        # Extract table contents if present
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    extracted_lines.append(" | ".join(row_cells))

        return "\n".join(extracted_lines)
    except Exception as e:
        logger.error("Error reading docx file %s: %s", file_path.name, e)
        return ""

def extract_from_pdf(file_path: Path) -> str:
    """
    Extract text from a PDF document using pdfplumber with automatic fallback
    to pypdf.PdfReader.

    Args:
        file_path: Path to the .pdf file.

    Returns:
        str: Extracted textual content.
    """
    extracted_text = ""

    # Primary parser: pdfplumber
    try:
        with pdfplumber.open(file_path) as pdf:
            pages_text = []
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text.strip())
            extracted_text = "\n".join(pages_text).strip()
    except Exception as e:
        logger.warning("pdfplumber extraction encountered error on %s: %s", file_path.name, e)
        extracted_text = ""

    # Fallback parser: pypdf if pdfplumber returned empty or encountered an error
    if not extracted_text:
        logger.info("Attempting pypdf fallback extraction for %s", file_path.name)
        try:
            reader = pypdf.PdfReader(str(file_path))
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text.strip())
            extracted_text = "\n".join(pages_text).strip()
        except Exception as e:
            logger.error("pypdf fallback also failed for %s: %s", file_path.name, e)
            extracted_text = ""

    return extracted_text

def extract_text(file_path: Path | str) -> str:
    """
    Master dispatcher extracting text from supported document types (.pdf, .docx, .txt).
    Gracefully returns an empty string for 0-byte, locked, or missing files.

    Args:
        file_path: File path as a string or Path object.

    Returns:
        str: Normalized, extracted text content.
    """
    path = Path(file_path)

    if not path.exists():
        logger.error("File does not exist: %s", path)
        return ""

    if path.is_file() and path.stat().st_size == 0:
        logger.warning("Empty (0-byte) file encountered: %s", path.name)
        return ""

    ext = path.suffix.lower()
    if ext == ".txt":
        return extract_from_txt(path)
    elif ext == ".docx":
        return extract_from_docx(path)
    elif ext == ".pdf":
        return extract_from_pdf(path)
    else:
        logger.warning("Unsupported file extension '%s' for: %s", ext, path.name)
        return ""
